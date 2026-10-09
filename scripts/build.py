"""Construit la base des lieux d'affectation SSA enrichie de critères de vie.

Entrée  : data/raw/sites_ssa.csv (liste des CMA / antennes médicales / HIA)
          + jeux de données publics téléchargés par scripts/download.py
Sorties : data/affectations.db      (SQLite, exploitable sous Python)
          output/affectations.xlsx  (tableau filtrable)
          output/sites.json         (données injectées dans la carte)

Usage : .venv/bin/python scripts/build.py
"""
import json
import math
import sqlite3
from pathlib import Path

import duckdb
import numpy as np
import pandas as pd

ROOT = Path(__file__).resolve().parent.parent
RAW = ROOT / "data" / "raw"
SIRENE = ROOT / "data" / "cache" / "sirene"
OUT = ROOT / "output"

# Rayons d'accès, à vol d'oiseau. 30 min de vélo ≈ 7,5 km de route à 15 km/h,
# soit ~6 km à vol d'oiseau ; 30 min de voiture ≈ 30 km de route ≈ 23 km.
R_VELO_KM = 6.0
R_VOITURE_KM = 23.0
FACTEUR_DETOUR = 1.3
SEUIL_GRANDE_VILLE = 80_000

EARTH_R = 6371.0

# Villes > 80 000 hab. absentes du référentiel communal (COM du Pacifique, étranger).
VILLES_HORS_REFERENTIEL = [
    # DROM sans commune > 80 000 hab. : on retient l'agglomération principale.
    ("Fort-de-France (agglomération)", 14.6161, -61.0588, 155_000),
    ("Pointe-à-Pitre (agglomération)", 16.2411, -61.5331, 240_000),
    ("Cayenne (agglomération)", 4.9372, -52.3260, 150_000),
    ("Mamoudzou (agglomération)", -12.7806, 45.2279, 110_000),
    ("Nouméa (agglomération)", -22.2758, 166.4580, 182_000),
    ("Papeete (agglomération)", -17.5350, -149.5696, 136_000),
    ("Djibouti", 11.5880, 43.1450, 600_000),
    ("Abidjan", 5.3600, -4.0083, 5_600_000),
    ("Dakar", 14.7167, -17.4677, 1_400_000),
    ("Libreville", 0.4162, 9.4673, 700_000),
    ("Abou Dabi", 24.4539, 54.3773, 1_500_000),
    ("Niamey", 13.5116, 2.1254, 1_300_000),
    ("N'Djamena", 12.1348, 15.0557, 1_500_000),
]

# Tranches d'effectif Sirene -> (libellé, effectif médian estimé)
TRANCHES = {
    "21": ("50-99", 75), "22": ("100-199", 150), "31": ("200-249", 225),
    "32": ("250-499", 375), "41": ("500-999", 750), "42": ("1000-1999", 1500),
    "51": ("2000-4999", 3500), "52": ("5000-9999", 7500), "53": ("10000+", 10000),
}

# Divisions NAF retenues pour l'emploi ingénieur -> (secteur, métiers d'ingénieur typiques)
SECTEURS = {
    "05": ("Industries extractives", "géologie, procédés, HSE"),
    "06": ("Industries extractives", "géologie, procédés, HSE"),
    "07": ("Industries extractives", "géologie, procédés, HSE"),
    "08": ("Industries extractives", "géologie, procédés, HSE"),
    "09": ("Industries extractives", "géologie, procédés, HSE"),
    "10": ("Agroalimentaire", "production, qualité, procédés, maintenance"),
    "11": ("Agroalimentaire", "production, qualité, procédés, maintenance"),
    "12": ("Agroalimentaire", "production, qualité, procédés"),
    "13": ("Textile / habillement", "production, matériaux, qualité"),
    "14": ("Textile / habillement", "production, matériaux, qualité"),
    "15": ("Textile / habillement", "production, matériaux, qualité"),
    "16": ("Bois / papier / imprimerie", "production, procédés, maintenance"),
    "17": ("Bois / papier / imprimerie", "production, procédés, maintenance"),
    "18": ("Bois / papier / imprimerie", "production, maintenance"),
    "19": ("Chimie / pétrochimie / pharmacie", "génie chimique, procédés, HSE, qualité"),
    "20": ("Chimie / pétrochimie / pharmacie", "génie chimique, procédés, HSE, qualité"),
    "21": ("Chimie / pétrochimie / pharmacie", "procédés pharma, qualité, validation"),
    "22": ("Plasturgie / matériaux", "matériaux, procédés, industrialisation"),
    "23": ("Plasturgie / matériaux", "matériaux, procédés, production"),
    "24": ("Métallurgie / mécanique", "mécanique, méthodes, production, qualité"),
    "25": ("Métallurgie / mécanique", "mécanique, méthodes, production, qualité"),
    "26": ("Électronique / informatique industrielle", "électronique, systèmes embarqués, test"),
    "27": ("Équipements électriques", "électrotechnique, conception, test"),
    "28": ("Machines et équipements", "conception mécanique, automatisme, méthodes"),
    "29": ("Automobile", "conception, industrialisation, qualité, supply chain"),
    "30": ("Aéronautique / naval / ferroviaire / défense", "aéronautique, systèmes, structures, MCO, essais"),
    "32": ("Industries diverses / dispositifs médicaux", "biomédical, qualité, industrialisation"),
    "33": ("Maintenance industrielle", "maintenance, fiabilité, MCO"),
    "35": ("Énergie", "énergie, électricité, nucléaire, réseaux"),
    "36": ("Eau / environnement", "traitement de l'eau, environnement"),
    "37": ("Eau / environnement", "environnement, assainissement"),
    "38": ("Eau / environnement", "déchets, valorisation, environnement"),
    "39": ("Eau / environnement", "dépollution, environnement"),
    "42": ("Génie civil / BTP", "génie civil, travaux, bureau d'études"),
    "61": ("Télécoms / numérique", "réseaux, télécoms, cybersécurité"),
    "62": ("Télécoms / numérique", "logiciel, data, cybersécurité, IT"),
    "63": ("Télécoms / numérique", "data, cloud, IT"),
    "71": ("Ingénierie / bureaux d'études / essais", "bureau d'études, conseil, essais, contrôle"),
    "72": ("R&D scientifique", "recherche, innovation, R&D"),
}


# ---------------------------------------------------------------- géométrie

def haversine(lat1, lon1, lat2, lon2):
    """Distance en km ; lat2/lon2 peuvent être des tableaux numpy."""
    lat1, lon1 = math.radians(lat1), math.radians(lon1)
    lat2, lon2 = np.radians(lat2), np.radians(lon2)
    a = np.sin((lat2 - lat1) / 2) ** 2 + math.cos(lat1) * np.cos(lat2) * np.sin((lon2 - lon1) / 2) ** 2
    return 2 * EARTH_R * np.arcsin(np.sqrt(a))


def minutes_voiture(km_vol):
    """Temps de trajet approximatif : distance à vol d'oiseau x détour, vitesse selon la distance."""
    if km_vol is None or (isinstance(km_vol, float) and math.isnan(km_vol)):
        return None
    route = km_vol * FACTEUR_DETOUR
    if route < 10:
        v = 40
    elif route < 40:
        v = 60
    else:
        v = 85
    return round(route / v * 60)


DVF_ABSENT = {"57", "67", "68", "976"}
REVERSE_CACHE = ROOT / "data" / "cache" / "reverse_communes.json"


def commune_at(lat, lon, communes):
    """Commune contenant le point (API Géo, mise en cache), sinon le centre de commune le plus proche."""
    import requests
    cache = json.loads(REVERSE_CACHE.read_text()) if REVERSE_CACHE.exists() else {}
    key = f"{lat:.4f},{lon:.4f}"
    if key not in cache:
        try:
            resp = requests.get("https://geo.api.gouv.fr/communes", params={"lat": lat, "lon": lon, "fields": "code"}, timeout=20)
            res = resp.json()
            cache[key] = res[0]["code"] if res else None
        except Exception:
            cache[key] = None
        REVERSE_CACHE.write_text(json.dumps(cache))
    code = cache[key]
    if code:
        hit = communes[communes.code == code]
        if not hit.empty:
            return hit.iloc[0]
    com, d = nearest(lat, lon, communes)
    return com if d is not None and d < 20 else None


def nearest(lat, lon, df):
    if df.empty:
        return None, None
    d = haversine(lat, lon, df.lat.values, df.lon.values)
    i = int(np.argmin(d))
    return df.iloc[i], float(d[i])


# ---------------------------------------------------------------- chargement

def load_sites():
    import os
    s = pd.read_csv(os.environ.get("SITES_CSV", RAW / "sites_ssa.csv"), dtype={"code_postal": str, "departement": str})
    s = s.dropna(subset=["latitude", "longitude"]).drop_duplicates("id")
    s["latitude"] = s.latitude.astype(float)
    s["longitude"] = s.longitude.astype(float)
    return s.reset_index(drop=True)


def load_communes():
    c = pd.DataFrame(json.load(open(RAW / "communes.json")))
    c = c[c.centre.notna()].copy()
    c["lon"] = c.centre.map(lambda g: g["coordinates"][0])
    c["lat"] = c.centre.map(lambda g: g["coordinates"][1])
    c["population"] = c.population.fillna(0).astype(int)
    return c[["code", "nom", "codeDepartement", "population", "lat", "lon"]]


def load_dvf():
    d = pd.read_csv(RAW / "dvf_stats.csv", dtype={"code_geo": str, "code_parent": str})
    d = d[d.echelle_geo == "commune"].copy()
    cols = ["nb_ventes_whole_appartement", "med_prix_m2_whole_appartement",
            "nb_ventes_whole_maison", "med_prix_m2_whole_maison",
            "nb_ventes_whole_apt_maison", "med_prix_m2_whole_apt_maison"]
    d = d[["code_geo"] + cols]
    # Paris, Lyon, Marseille : DVF est par arrondissement, le référentiel par commune.
    plm = {"75056": "751", "69123": "6938", "13055": "132"}
    extra = []
    for code, prefix in plm.items():
        sub = d[d.code_geo.str.startswith(prefix) & (d.code_geo.str.len() == 5)]
        if sub.empty:
            continue
        row = {"code_geo": code}
        for t in ["appartement", "maison", "apt_maison"]:
            n, p = sub[f"nb_ventes_whole_{t}"].fillna(0), sub[f"med_prix_m2_whole_{t}"]
            ok = p.notna() & (n > 0)
            row[f"nb_ventes_whole_{t}"] = n.sum()
            row[f"med_prix_m2_whole_{t}"] = (p[ok] * n[ok]).sum() / n[ok].sum() if ok.any() else np.nan
        extra.append(row)
    d = pd.concat([d, pd.DataFrame(extra)], ignore_index=True)
    return d.drop_duplicates("code_geo").set_index("code_geo")


def load_tgv():
    st = pd.read_csv(RAW / "gtfs_tgv" / "stops.txt", dtype=str)
    st = st[st.location_type == "1"].copy()
    st["lat"], st["lon"] = st.stop_lat.astype(float), st.stop_lon.astype(float)
    st["nom"] = st.stop_name.str.title()
    # Gares françaises uniquement (le flux contient aussi des gares étrangères desservies)
    st = st[st.stop_id.str.contains("OCE87")]
    return st[["nom", "lat", "lon"]].drop_duplicates("nom").reset_index(drop=True)


def load_airports():
    a = pd.read_csv(RAW / "airports.csv", low_memory=False)
    a = a[(a.scheduled_service == "yes") & a.type.isin(["large_airport", "medium_airport"])].copy()
    a = a.rename(columns={"latitude_deg": "lat", "longitude_deg": "lon", "name": "nom"})
    a["code"] = a.iata_code.fillna(a.ident)
    return a[["nom", "code", "type", "iso_country", "municipality", "lat", "lon"]].reset_index(drop=True)


def load_big_cities(communes):
    big = communes[communes.population >= SEUIL_GRANDE_VILLE][["nom", "lat", "lon", "population"]].copy()
    extra = pd.DataFrame(VILLES_HORS_REFERENTIEL, columns=["nom", "lat", "lon", "population"])
    return pd.concat([big, extra], ignore_index=True)


def load_culture():
    b = pd.read_csv(RAW / "basilic.csv", sep=";", low_memory=False,
                    usecols=["Nom", "Type équipement ou lieu", "Latitude", "Longitude",
                             "Nombre_ecrans", "libelle_geographique"])
    b = b.rename(columns={"Type équipement ou lieu": "type", "Latitude": "lat", "Longitude": "lon",
                          "Nombre_ecrans": "ecrans", "libelle_geographique": "commune"})
    b = b.dropna(subset=["lat", "lon"])
    groupes = {
        "Cinéma": "cinema",
        "Théâtre": "spectacle", "Scène": "spectacle", "Opéra": "spectacle",
        "Centre de création musicale": "spectacle",
        "Musée": "musee", "Centre d'art": "musee", "Centre de création artistique": "musee",
        "Bibliothèque": "bibliotheque",
        "Conservatoire": "conservatoire",
        "Librairie": "librairie",
        "Monument": "monument",
    }
    b["groupe"] = b.type.map(groupes)
    b = b[b.groupe.notna()].reset_index(drop=True)
    b["ecrans"] = pd.to_numeric(b.ecrans, errors="coerce").fillna(0)
    return b


def load_festivals():
    f = pd.read_csv(RAW / "festivals.csv", sep=";", usecols=["Nom de la manifestation", "Domaine", "coordonnees_insee"])
    f = f.dropna(subset=["coordonnees_insee"])
    ll = f.coordonnees_insee.str.split(",", expand=True).astype(float)
    f["lat"], f["lon"] = ll[0], ll[1]
    return f.rename(columns={"Nom de la manifestation": "nom", "Domaine": "domaine"})


def load_education():
    e = pd.read_csv(RAW / "education_plain.csv", sep=";", dtype=str,
                    usecols=["type_etablissement", "etat", "ecole_maternelle", "ecole_elementaire",
                             "latitude", "longitude", "statut_public_prive"])
    e = e[e.etat == "OUVERT"].dropna(subset=["latitude", "longitude"]).copy()
    e["lat"], e["lon"] = e.latitude.astype(float), e.longitude.astype(float)
    e["maternelle"] = e.ecole_maternelle == "1"
    e["elementaire"] = e.ecole_elementaire == "1"
    e["college"] = e.type_etablissement == "Collège"
    e["lycee"] = e.type_etablissement == "Lycée"
    return e[["lat", "lon", "maternelle", "elementaire", "college", "lycee"]].reset_index(drop=True)


def load_sirene():
    """Crèches (NAF 88.91A) et employeurs >= 50 salariés des secteurs ingénieur, géolocalisés."""
    cache = SIRENE / "extrait.parquet"
    if cache.exists():
        return pd.read_parquet(cache)
    divs = ",".join(f"'{d}'" for d in SECTEURS)
    tranches = ",".join(f"'{t}'" for t in TRANCHES)
    con = duckdb.connect()
    df = con.sql(f"""
        with e as (
          select siret, siren, activitePrincipaleEtablissement as naf,
                 trancheEffectifsEtablissement as tranche,
                 coalesce(enseigne1Etablissement, denominationUsuelleEtablissement) as enseigne,
                 libelleCommuneEtablissement as commune
          from '{SIRENE / "etab.parquet"}'
          where etatAdministratifEtablissement = 'A'
            and ( activitePrincipaleEtablissement = '88.91A'
                  or (substr(activitePrincipaleEtablissement,1,2) in ({divs})
                      and trancheEffectifsEtablissement in ({tranches})) )
        )
        select e.*, u.denominationUniteLegale as raison_sociale, u.sigleUniteLegale as sigle,
               g.y_latitude as lat, g.x_longitude as lon
        from e
        join '{SIRENE / "geoloc.parquet"}' g using (siret)
        left join '{SIRENE / "ul.parquet"}' u using (siren)
        where g.y_latitude is not null
    """).df()
    df["creche"] = df.naf == "88.91A"
    df.to_parquet(cache)
    return df


# ---------------------------------------------------------------- calcul

def within(lat, lon, df, r):
    d = haversine(lat, lon, df.lat.values, df.lon.values)
    return df[d <= r], d[d <= r]


def nom_employeur(row):
    base = row.raison_sociale if isinstance(row.raison_sociale, str) and row.raison_sociale else row.enseigne
    if not isinstance(base, str) or not base:
        base = "(nom non diffusé)"
    if isinstance(row.enseigne, str) and row.enseigne and row.enseigne != base:
        base = f"{base} – {row.enseigne}"
    return base.title() if base.isupper() else base


def enrich(sites, communes, dvf, tgv, airports, intl_airports, big, culture, festivals, edu, sirene):
    creches = sirene[sirene.creche].reset_index(drop=True)
    emp = sirene[~sirene.creche].reset_index(drop=True)
    emp["div"] = emp.naf.str[:2]
    emp["secteur"] = emp["div"].map(lambda d: SECTEURS[d][0])
    emp["effectif_est"] = emp.tranche.map(lambda t: TRANCHES[t][1])
    emp["tranche_lib"] = emp.tranche.map(lambda t: TRANCHES[t][0])

    rows, employeurs = [], []
    for s in sites.itertuples():
        lat, lon = s.latitude, s.longitude
        r = {"id": s.id}
        france = s.zone != "Étranger"

        # Commune d'implantation et population des zones
        com = commune_at(lat, lon, communes) if france else None
        if com is not None:
            r["commune_insee"] = com.code
            r["commune_ref"] = com.nom
            r["population_commune"] = int(com.population)
            zv, _ = within(lat, lon, communes, R_VELO_KM)
            zc, _ = within(lat, lon, communes, R_VOITURE_KM)
            # Une grande commune peut avoir son centre hors du cercle : on l'inclut toujours.
            if com.code not in set(zv.code):
                zv = pd.concat([zv, communes[communes.code == com.code]])
            if com.code not in set(zc.code):
                zc = pd.concat([zc, communes[communes.code == com.code]])
            r["population_zone_velo"] = int(zv.population.sum())
            r["population_zone_voiture"] = int(zc.population.sum())
            # Prix immobiliers (DVF, médiane des ventes). DVF ne couvre pas l'Alsace-Moselle
            # ni Mayotte : on n'affiche pas une moyenne qui ne porterait que sur les communes voisines.
            if com.code[:2] in DVF_ABSENT or com.code[:3] in DVF_ABSENT:
                zv = zv.iloc[0:0]
                zc = zc.iloc[0:0]
            p = dvf.loc[com.code] if com.code in dvf.index else None
            r["prix_m2_appart_commune"] = None if p is None or pd.isna(p.med_prix_m2_whole_appartement) else round(p.med_prix_m2_whole_appartement)
            r["prix_m2_maison_commune"] = None if p is None or pd.isna(p.med_prix_m2_whole_maison) else round(p.med_prix_m2_whole_maison)
            for label, z in (("velo", zv), ("voiture", zc)):
                sub = dvf.reindex(z.code).dropna(subset=["med_prix_m2_whole_apt_maison"])
                n = sub.nb_ventes_whole_apt_maison.fillna(0)
                r[f"prix_m2_zone_{label}"] = round((sub.med_prix_m2_whole_apt_maison * n).sum() / n.sum()) if n.sum() > 0 else None
        else:
            for k in ["commune_insee", "commune_ref", "population_commune", "population_zone_velo",
                      "population_zone_voiture", "prix_m2_appart_commune", "prix_m2_maison_commune",
                      "prix_m2_zone_velo", "prix_m2_zone_voiture"]:
                r[k] = None

        # Grande ville, TGV, aéroports
        v, dv = nearest(lat, lon, big)
        r["grande_ville"] = v.nom
        r["grande_ville_pop"] = int(v.population)
        r["grande_ville_km"] = round(dv, 1)
        r["grande_ville_min"] = minutes_voiture(dv)
        r["grande_ville_lat"], r["grande_ville_lon"] = float(v.lat), float(v.lon)

        if s.zone == "Métropole" and str(s.departement) not in ("2A", "2B", "20"):
            g, dg = nearest(lat, lon, tgv)
            r["gare_tgv"] = g.nom
            r["gare_tgv_km"] = round(dg, 1)
            r["gare_tgv_min"] = minutes_voiture(dg)
            r["gare_tgv_lat"], r["gare_tgv_lon"] = float(g.lat), float(g.lon)
        else:
            r.update(gare_tgv=None, gare_tgv_km=None, gare_tgv_min=None, gare_tgv_lat=None, gare_tgv_lon=None)

        a, da = nearest(lat, lon, airports)
        r["aeroport"] = f"{a.nom} ({a.code})"
        r["aeroport_km"] = round(da, 1)
        r["aeroport_min"] = minutes_voiture(da)
        r["aeroport_lat"], r["aeroport_lon"] = float(a.lat), float(a.lon)
        ai, dai = nearest(lat, lon, intl_airports)
        r["aeroport_international"] = f"{ai.nom} ({ai.code})"
        r["aeroport_international_km"] = round(dai, 1)
        r["aeroport_international_min"] = minutes_voiture(dai)

        # Culture
        for label, rad in (("velo", R_VELO_KM), ("voiture", R_VOITURE_KM)):
            c, _ = within(lat, lon, culture, rad)
            g = c.groupby("groupe")
            counts = g.size()
            for grp in ["cinema", "spectacle", "musee", "bibliotheque", "conservatoire", "librairie", "monument"]:
                r[f"{grp}_{label}"] = int(counts.get(grp, 0))
            r[f"ecrans_cinema_{label}"] = int(c[c.groupe == "cinema"].ecrans.sum())
            f, _ = within(lat, lon, festivals, rad)
            r[f"festivals_{label}"] = len(f)

        # Petite enfance et écoles
        for label, rad in (("velo", R_VELO_KM), ("voiture", R_VOITURE_KM)):
            e, _ = within(lat, lon, edu, rad)
            r[f"ecoles_maternelles_{label}"] = int(e.maternelle.sum())
            r[f"ecoles_elementaires_{label}"] = int(e.elementaire.sum())
            r[f"colleges_{label}"] = int(e.college.sum())
            r[f"lycees_{label}"] = int(e.lycee.sum())
            cr, _ = within(lat, lon, creches, rad)
            r[f"creches_{label}"] = len(cr)

        # Industrie et emploi ingénieur (établissements >= 50 salariés)
        ez, dz = within(lat, lon, emp, R_VOITURE_KM)
        ez = ez.assign(dist=dz)
        r["employeurs_50plus_velo"] = int((dz <= R_VELO_KM).sum())
        r["employeurs_50plus_voiture"] = len(ez)
        r["emplois_estimes_voiture"] = int(ez.effectif_est.sum())
        by_sect = ez.groupby("secteur").effectif_est.sum().sort_values(ascending=False)
        top = list(by_sect.index[:3])
        r["secteur_dominant"] = top[0] if top else None
        r["secteurs_principaux"] = " ; ".join(f"{k} ({int(by_sect[k]):,} empl.)".replace(",", " ") for k in top) or None
        metiers = []
        for k in top:
            for d, (sect, met) in SECTEURS.items():
                if sect == k:
                    metiers.append(f"{k} : {met}")
                    break
        r["metiers_ingenieur"] = " | ".join(metiers) or None
        n_aero = ez[ez["div"] == "30"]
        r["employeurs_aero_defense"] = len(n_aero)
        best = ez.sort_values(["effectif_est", "dist"], ascending=[False, True])
        best = best.drop_duplicates("siren").head(12)
        for rank, e in enumerate(best.itertuples(), 1):
            employeurs.append({
                "site_id": s.id, "rang": rank, "nom": nom_employeur(e), "commune": (e.commune or "").title(),
                "effectif": e.tranche_lib, "secteur": e.secteur, "naf": e.naf,
                "distance_km": round(float(e.dist), 1), "lat": float(e.lat), "lon": float(e.lon),
            })
        r["principaux_employeurs"] = " ; ".join([f"{e['nom']} ({e['effectif']})" for e in employeurs if e["site_id"] == s.id][:6]) or None
        rows.append(r)

    return sites.merge(pd.DataFrame(rows), on="id"), pd.DataFrame(employeurs)


# ---------------------------------------------------------------- export

COLONNES_EXCEL = [
    ("nom", "Lieu d'affectation"), ("type_structure", "Type"), ("armee_soutenue", "Armée soutenue"),
    ("unites_soutenues", "Unités soutenues"), ("cma_rattachement", "CMA de rattachement"),
    ("commune", "Commune"), ("departement", "Dépt"), ("zone", "Zone"),
    ("population_commune", "Habitants commune"), ("population_zone_velo", "Habitants zone vélo (30 min)"),
    ("population_zone_voiture", "Habitants zone voiture (30 min)"),
    ("grande_ville", "Grande ville (>80k) la plus proche"), ("grande_ville_km", "Grande ville km"),
    ("grande_ville_min", "Grande ville min (voiture)"),
    ("gare_tgv", "Gare TGV la plus proche"), ("gare_tgv_km", "Gare TGV km"), ("gare_tgv_min", "Gare TGV min"),
    ("aeroport", "Aéroport le plus proche"), ("aeroport_km", "Aéroport km"), ("aeroport_min", "Aéroport min"),
    ("aeroport_international", "Grand aéroport international"), ("aeroport_international_min", "Aéroport intl min"),
    ("prix_m2_appart_commune", "Prix m² appart. commune (€)"), ("prix_m2_maison_commune", "Prix m² maison commune (€)"),
    ("prix_m2_zone_velo", "Prix m² moyen zone vélo (€)"), ("prix_m2_zone_voiture", "Prix m² moyen zone voiture (€)"),
    ("creches_velo", "Crèches zone vélo"), ("creches_voiture", "Crèches zone voiture"),
    ("ecoles_maternelles_velo", "Écoles maternelles zone vélo"), ("ecoles_elementaires_velo", "Écoles élémentaires zone vélo"),
    ("ecoles_maternelles_voiture", "Écoles maternelles zone voiture"), ("ecoles_elementaires_voiture", "Écoles élémentaires zone voiture"),
    ("colleges_voiture", "Collèges zone voiture"), ("lycees_voiture", "Lycées zone voiture"),
    ("cinema_velo", "Cinémas zone vélo"), ("cinema_voiture", "Cinémas zone voiture"), ("ecrans_cinema_voiture", "Écrans de cinéma zone voiture"),
    ("spectacle_velo", "Salles de spectacle zone vélo"), ("spectacle_voiture", "Salles de spectacle zone voiture"),
    ("musee_voiture", "Musées / centres d'art zone voiture"), ("bibliotheque_voiture", "Bibliothèques zone voiture"),
    ("conservatoire_voiture", "Conservatoires zone voiture"), ("librairie_voiture", "Librairies zone voiture"),
    ("festivals_voiture", "Festivals zone voiture"), ("monument_voiture", "Monuments historiques zone voiture"),
    ("employeurs_50plus_velo", "Employeurs ingé ≥50 sal. zone vélo"), ("employeurs_50plus_voiture", "Employeurs ingé ≥50 sal. zone voiture"),
    ("emplois_estimes_voiture", "Emplois estimés (ces employeurs)"), ("employeurs_aero_defense", "Sites aéro/naval/défense ≥50 sal."),
    ("secteur_dominant", "Secteur industriel dominant"), ("secteurs_principaux", "3 secteurs principaux"),
    ("metiers_ingenieur", "Métiers d'ingénieur typiques"), ("principaux_employeurs", "Principaux employeurs"),
    ("source_confiance", "Fiabilité de la ligne"), ("source_url", "Source"), ("remarques", "Remarques"),
    ("latitude", "Latitude"), ("longitude", "Longitude"),
]


def export_excel(df, emp, path):
    from openpyxl.formatting.rule import ColorScaleRule
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    cols = [c for c, _ in COLONNES_EXCEL if c in df.columns]
    t = df[cols].rename(columns=dict(COLONNES_EXCEL))
    t.insert(0, "Ma sélection (x)", "")
    t.insert(1, "Notes", "")
    t = t.sort_values(["Zone", "Armée soutenue", "Lieu d'affectation"])

    e = emp.merge(df[["id", "nom"]], left_on="site_id", right_on="id").drop(columns=["id", "site_id", "lat", "lon"])
    e = e.rename(columns={"nom_x": "Employeur", "nom_y": "Lieu d'affectation", "rang": "Rang", "commune": "Commune",
                          "effectif": "Effectif", "secteur": "Secteur", "naf": "Code NAF", "distance_km": "Distance (km)"})
    e = e[["Lieu d'affectation", "Rang", "Employeur", "Secteur", "Effectif", "Commune", "Distance (km)", "Code NAF"]]

    methode = pd.DataFrame([
        ("Lieux d'affectation", "CMA, antennes médicales et HIA du Service de santé des armées, compilés depuis defense.gouv.fr et sources publiques. 'déduite' = antenne supposée car une unité y stationne, à confirmer."),
        ("Zone vélo", f"Cercle de {R_VELO_KM:g} km à vol d'oiseau (≈ 30 min à 15 km/h avec détours)."),
        ("Zone voiture", f"Cercle de {R_VOITURE_KM:g} km à vol d'oiseau (≈ 30 min de route hors bouchons)."),
        ("Temps de trajet", f"Estimation : distance à vol d'oiseau x {FACTEUR_DETOUR} puis 40 / 60 / 85 km/h selon la distance. Ordre de grandeur, pas un itinéraire."),
        ("Habitants", "Population municipale Insee (geo.api.gouv.fr) ; zones = somme des communes dont le centre est dans le cercle."),
        ("Grande ville", f"Commune de plus de {SEUIL_GRANDE_VILLE:,} habitants la plus proche.".replace(",", " ")),
        ("TGV", "Gares desservies par TGV INOUI / OUIGO (GTFS SNCF Voyageurs)."),
        ("Aéroports", "OurAirports : aéroports avec vols réguliers. 'International' = grand aéroport (catégorie large)."),
        ("Prix m²", "DVF (Etalab, ventes 5 dernières années), médiane par commune ; zone = moyenne pondérée par le nombre de ventes. Non disponible en Alsace-Moselle, à Mayotte et hors DROM."),
        ("Crèches", "Établissements actifs Sirene en NAF 88.91A (accueil de jeunes enfants). Les crèches gérées directement par une mairie ne sont pas toujours comptées."),
        ("Écoles", "Annuaire de l'Éducation nationale, établissements ouverts."),
        ("Culture", "Base Basilic (ministère de la Culture) + Panorama des festivals."),
        ("Industrie", "Établissements Sirene actifs ≥ 50 salariés dans l'industrie, l'énergie, le génie civil, le numérique, l'ingénierie et la R&D. Emplois estimés = milieu de la tranche d'effectif."),
    ], columns=["Critère", "Méthode et source"])

    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        t.to_excel(xw, sheet_name="Lieux d'affectation", index=False)
        e.to_excel(xw, sheet_name="Employeurs proches", index=False)
        methode.to_excel(xw, sheet_name="Méthode et sources", index=False)

        head_fill = PatternFill("solid", fgColor="1F3A4D")
        for ws in xw.book.worksheets:
            ws.freeze_panes = "C2" if ws.title == "Lieux d'affectation" else "A2"
            ws.auto_filter.ref = ws.dimensions
            for cell in ws[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = head_fill
                cell.alignment = Alignment(wrap_text=True, vertical="center")
            ws.row_dimensions[1].height = 48
            for i, col in enumerate(ws.columns, 1):
                width = max(len(str(c.value or "")) for c in list(col)[:200])
                ws.column_dimensions[get_column_letter(i)].width = min(max(10, width * 0.9), 60)

        ws = xw.book["Lieux d'affectation"]
        hdr = {c.value: c.column_letter for c in ws[1]}
        n = ws.max_row
        good_high = ["Habitants zone vélo (30 min)", "Crèches zone vélo", "Écoles maternelles zone vélo",
                     "Cinémas zone voiture", "Salles de spectacle zone voiture", "Festivals zone voiture",
                     "Employeurs ingé ≥50 sal. zone voiture"]
        good_low = ["Gare TGV min", "Aéroport min", "Grande ville min (voiture)", "Prix m² moyen zone voiture (€)"]
        for name in good_high:
            if name in hdr:
                ws.conditional_formatting.add(f"{hdr[name]}2:{hdr[name]}{n}", ColorScaleRule(
                    start_type="percentile", start_value=5, start_color="F8D7C4",
                    end_type="percentile", end_value=95, end_color="B9DFC3"))
        for name in good_low:
            if name in hdr:
                ws.conditional_formatting.add(f"{hdr[name]}2:{hdr[name]}{n}", ColorScaleRule(
                    start_type="percentile", start_value=5, start_color="B9DFC3",
                    end_type="percentile", end_value=95, end_color="F8D7C4"))
        ws.column_dimensions["A"].width = 12
        ws.column_dimensions["B"].width = 24
        ws.column_dimensions["C"].width = 42


def export_sqlite(df, emp, refs, path):
    if path.exists():
        path.unlink()
    con = sqlite3.connect(path)
    df.to_sql("sites", con, index=False)
    emp.to_sql("employeurs", con, index=False)
    for name, t in refs.items():
        t.to_sql(name, con, index=False)
    con.execute("create index idx_emp_site on employeurs(site_id)")
    con.execute("""create view v_synthese as
        select nom, type_structure, armee_soutenue, commune, departement, zone,
               population_zone_velo, grande_ville, grande_ville_min, gare_tgv, gare_tgv_min,
               aeroport, aeroport_min, prix_m2_zone_voiture, creches_velo, ecoles_maternelles_velo,
               cinema_voiture, festivals_voiture, employeurs_50plus_voiture, secteur_dominant
        from sites""")
    con.commit()
    con.close()


def export_json(df, emp, culture, path):
    keep_emp = emp.drop(columns=["naf"]).to_dict("records")
    by_site = {}
    for e in keep_emp:
        by_site.setdefault(e.pop("site_id"), []).append(e)
    recs = json.loads(df.to_json(orient="records", force_ascii=False))
    for r in recs:
        r["employeurs"] = by_site.get(r["id"], [])

    # Points culturels affichés sur la carte (cinémas, scènes, musées) proches d'au moins un site
    poi = culture[culture.groupe.isin(["cinema", "spectacle", "musee"])].copy()
    keep = np.zeros(len(poi), dtype=bool)
    for s in df.itertuples():
        keep |= haversine(s.latitude, s.longitude, poi.lat.values, poi.lon.values) <= R_VOITURE_KM
    poi = poi[keep]
    pois = [[round(p.lat, 5), round(p.lon, 5), p.groupe[0], p.Nom] for p in poi.itertuples()]

    payload = {"sites": recs, "poi": pois,
               "rayons": {"velo": R_VELO_KM, "voiture": R_VOITURE_KM}}
    path.write_text(json.dumps(payload, ensure_ascii=False, separators=(",", ":")))


def main():
    OUT.mkdir(exist_ok=True)
    print("chargement…")
    sites = load_sites()
    communes = load_communes()
    dvf = load_dvf()
    tgv = load_tgv()
    airports = load_airports()
    intl = airports[airports.type == "large_airport"].reset_index(drop=True)
    big = load_big_cities(communes)
    culture = load_culture()
    festivals = load_festivals()
    edu = load_education()
    sirene = load_sirene()
    print(f"{len(sites)} sites, {len(sirene)} établissements Sirene retenus")

    df, emp = enrich(sites, communes, dvf, tgv, airports, intl, big, culture, festivals, edu, sirene)

    export_sqlite(df, emp, {"ref_gares_tgv": tgv, "ref_aeroports": airports, "ref_grandes_villes": big},
                  ROOT / "data" / "affectations.db")
    export_excel(df, emp, OUT / "affectations.xlsx")
    export_json(df, emp, culture, OUT / "sites.json")
    print(f"ok : {len(df)} lieux, {len(emp)} lignes employeurs")


if __name__ == "__main__":
    main()
